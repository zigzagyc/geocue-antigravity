from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

def create_presentation():
    prs = Presentation()

    # Define a simple clean theme by manually setting elements (since we don't have a template)
    # We will use a standard slide layout for simplicity and focus on content

    # -------------------------------------------------------------------------
    # Slide 1: Title Slide
    # -------------------------------------------------------------------------
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]

    title.text = "HearHere"
    subtitle.text = "AI-Powered Location Accessibility & Community Engagement\n\nCreators: Randall Zhang, Runxin Tao, Jefferson Liu, Xiyao Sha\nCoach: Yuecheng Zhang\n\n2026 Presidential AI Challenge - Track II"

    # Adjust title formatting
    # title.text_frame.paragraphs[0].font.size = Pt(60)
    # title.text_frame.paragraphs[0].font.bold = True
    # title.text_frame.paragraphs[0].font.color.rgb = RGBColor(0, 51, 102) # Dark Blue

    # -------------------------------------------------------------------------
    # Slide 2: The Challenge
    # -------------------------------------------------------------------------
    bullet_slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]

    title_shape.text = "The Invisible Barriers"

    tf = body_shape.text_frame
    tf.text = "Visual Impairment"
    p = tf.paragraphs[0]
    p.font.size = Pt(24)
    p.font.bold = True
    
    p = tf.add_paragraph()
    p.text = "Over 2.2 Billion people globally face vision challenges. Static signs don't talk back."
    p.level = 1

    p = tf.add_paragraph()
    p.text = "Language Barriers"
    p.font.bold = True
    p.font.size = Pt(24)

    p = tf.add_paragraph()
    p.text = "Local history and safety info are often English-only, isolating non-native speakers."
    p.level = 1

    p = tf.add_paragraph()
    p.text = "Static Information"
    p.font.bold = True
    p.font.size = Pt(24)

    p = tf.add_paragraph()
    p.text = "Physical signs are expensive to update and never adapt to the user."
    p.level = 1

    # -------------------------------------------------------------------------
    # Slide 3: The Solution - HearHere
    # -------------------------------------------------------------------------
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]

    title_shape.text = "Introducing HearHere"

    tf = body_shape.text_frame
    tf.text = "Anchoring Digital Knowledge to Physical Spaces"
    
    p = tf.add_paragraph()
    p.text = "An invisible layer of accessibility over the real world."
    p.level = 1

    p = tf.add_paragraph()
    p.text = "\"Cues\": Location-based audio markers."
    p.level = 1

    p = tf.add_paragraph()
    p.text = "Hands-Free: Information plays automatically when you are nearby."
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "Inclusive: For the visually impaired, the curious tourist, and the local historian."
    p.level = 1

    # -------------------------------------------------------------------------
    # Slide 4: AI Power - Gemini 1.5 Flash
    # -------------------------------------------------------------------------
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]

    title_shape.text = "Powered by Google Gemini 1.5 Flash"

    tf = body_shape.text_frame
    tf.text = "Why AI? To democratize accessibility creation."

    p = tf.add_paragraph()
    p.text = "Intelligent Speech-to-Text (STT)"
    p.font.bold = True
    
    p = tf.add_paragraph()
    p.text = "Create Cues just by speaking. AI cleans up the text for you."
    p.level = 1

    p = tf.add_paragraph()
    p.text = "Generative Translation"
    p.font.bold = True

    p = tf.add_paragraph()
    p.text = "Breaking language barriers in real-time."
    p.level = 1

    p = tf.add_paragraph()
    p.text = "Context-aware translation (preserves nuance over literal translation)."
    p.level = 1

    # -------------------------------------------------------------------------
    # Slide 5: Key Features
    # -------------------------------------------------------------------------
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]

    title_shape.text = "Key Features"

    tf = body_shape.text_frame
    tf.text = "Proximity-Based Playback"
    tf.paragraphs[0].font.bold = True

    p = tf.add_paragraph()
    p.text = "Safety first. No need to look at the screen while walking."
    p.level = 1

    p = tf.add_paragraph()
    p.text = "Universal Creation"
    p.font.bold = True
    
    p = tf.add_paragraph()
    p.text = "Anyone can leave a Cue. From \"Wet Cement\" warnings to historic facts."
    p.level = 1

    p = tf.add_paragraph()
    p.text = "Casting Support"
    p.font.bold = True
    
    p = tf.add_paragraph()
    p.text = "Project audio to Google Home/AirPlay for classroom or group tour settings."
    p.level = 1

    # -------------------------------------------------------------------------
    # Slide 6: Technical Architecture
    # -------------------------------------------------------------------------
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]

    title_shape.text = "Technical Architecture"

    tf = body_shape.text_frame
    tf.text = "Frontend: Flutter (Dart) - Single codebase for iOS and Android."
    
    p = tf.add_paragraph()
    p.text = "State Management: Riverpod for robustness."
    
    p = tf.add_paragraph()
    p.text = "AI Service: Google Gemini API integration."
    
    p = tf.add_paragraph()
    p.text = "Background Efficiency: ProximityService smart location monitoring."

    # Note: Adding the diagram as text description or placeholder logic
    # In a real scenario, you'd insert an image of the diagram here
    p = tf.add_paragraph()
    p.text = "[Architecture Diagram Placeholder]"
    p.level = 1

    # -------------------------------------------------------------------------
    # Slide 7: Future Roadmap
    # -------------------------------------------------------------------------
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]

    title_shape.text = "Future Roadmap"

    tf = body_shape.text_frame
    tf.text = "Short Term: Computer Vision"
    tf.paragraphs[0].font.bold = True

    p = tf.add_paragraph()
    p.text = "\"Snap to Ask\" - Take a photo of a building to generate a Cue."
    p.level = 1

    p = tf.add_paragraph()
    p.text = "Mid Term: Personalized Tours"
    p.font.bold = True
    
    p = tf.add_paragraph()
    p.text = "\"Show me 19th-century architecture near me.\""
    p.level = 1

    p = tf.add_paragraph()
    p.text = "Long Term: Hazard Detection"
    p.font.bold = True
    
    p = tf.add_paragraph()
    p.text = "Aggregating user Cues to help municipalities fix unsafe areas."
    p.level = 1

    # -------------------------------------------------------------------------
    # Slide 8: Conclusion
    # -------------------------------------------------------------------------
    slide = prs.slides.add_slide(bullet_slide_layout)
    shapes = slide.shapes
    title_shape = shapes.title
    body_shape = shapes.placeholders[1]

    title_shape.text = "Conclusion"

    tf = body_shape.text_frame
    tf.text = "A Digital Layer of Empathy"
    tf.paragraphs[0].font.bold = True
    
    p = tf.add_paragraph()
    p.text = "HearHere re-connects us with our neighborhoods and makes \"Accessibility\" a community effort."
    p.level = 1

    p = tf.add_paragraph()
    p.text = "Call to Action"
    p.font.bold = True
    
    p = tf.add_paragraph()
    p.text = "Let's make every location speak to everyone."
    p.level = 1

    # Save
    output_path = 'docs/HearHere_Presentation.pptx'
    prs.save(output_path)
    print(f"Presentation saved to {output_path}")

if __name__ == "__main__":
    create_presentation()
